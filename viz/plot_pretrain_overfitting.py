"""Show pretraining overfitting in the pretrain-epoch sweep (seed 42, finetune on S1000D_spec_sample).

Reads the HF Trainer lines in runs/xml/<tag>_logs/<tag>_<source>_to_S1000D_spec_sample.pretrain.log:
  {'loss': ..., 'epoch': ...}       one per step (logging_steps=1)
  {'eval_loss': ..., 'epoch': ...}  one per 5 steps on the source's pretrain dev split (eval_steps=5)

1. Writes results/figures/pretrain_overfitting.{pdf,png}:
   top: train loss (per-epoch mean) and dev loss over the 100-epoch run, one panel per source;
   bottom: final pretrain dev loss and test MRR against pretrain epochs.
2. Updates results/05_ablation_study/pretrain_epoch_sweep.xlsx:
   - "Final pretrain dev loss" columns (last eval_loss of each run) after the final step loss columns;
   - sheets "Figure" (the PNG + findings), "Pretrain losses" (one row per run and source) and
     "Dev loss curves" (the 100-epoch curves plotted in the figure);
   - Notes rows describing them.
   Re-running overwrites all of these in place. openpyxl drops cached formula values, so recalc afterwards
   (e.g. soffice --headless --calc --convert-to xlsx --outdir <tmp> <xlsx>, then copy back).

Run with .venv/bin/python viz/plot_pretrain_overfitting.py
"""
import json
import re
from collections import defaultdict
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.drawing.image import Image
from matplotlib.ticker import FixedLocator, NullLocator
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
RUNS = REPO / "runs" / "xml"
XLSX = REPO / "results" / "05_ablation_study" / "pretrain_epoch_sweep.xlsx"
OUT = REPO / "results" / "figures"
TARGET = "S1000D_spec_sample"
SOURCES = ["BIKE_7", "FOSSIG", "s1kd-tools-doc"]
EPOCHS = [5, 10, 30, 50, 100]

# Run tag behind each Excel row (pretrain epochs, setup). Original row 9 (README Ablation 2, 10 ep) has no logs.
SHEETS = {
    "Original queries": {
        "tag": lambda ep: f"longpretrain{ep}_s42",
        "results": REPO / "results" / "05_ablation_study",
        "rows": {5: "longpretrain5_s42", 6: "longpretrain10_s42", 7: "longpretrain30_s42",
                 8: "longpretrain50_s42", 9: None, 10: "longpretrain100_s42"},
    },
    "Anonymized queries": {
        "tag": lambda ep: f"anonq1024_longpretrain{ep}_s42",
        "results": REPO / "results" / "07_anonymized",
        "rows": {5: "anonq1024_longpretrain5_s42", 6: "anonq1024_longpretrain10_s42",
                 7: "anonq1024_longpretrain30_s42", 8: "anonq1024_longpretrain50_s42",
                 9: "anonq1024_meancos_p1024_s42", 10: "anonq1024_longpretrain100_s42"},
    },
}

# Reference categorical palette, first three slots (validated all-pairs); markers add a second encoding.
COLORS = {"BIKE_7": "#2a78d6", "FOSSIG": "#eb6834", "s1kd-tools-doc": "#1baf7a"}
MARKERS = {"BIKE_7": "o", "FOSSIG": "s", "s1kd-tools-doc": "^"}
INK, MUTED, GRID = "#0b0b0b", "#52514e", "#e4e3df"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5, "axes.edgecolor": MUTED,
                     "axes.labelcolor": INK, "xtick.color": MUTED, "ytick.color": MUTED})

LOSS_RE = re.compile(r"\{'loss': ([0-9.]+), 'learning_rate': [^,]+, 'epoch': ([0-9.]+)\}")
EVAL_RE = re.compile(r"\{'eval_loss': ([0-9.]+),[^}]*'epoch': ([0-9.]+)\}")


def parse_log(tag, source):
    path = RUNS / f"{tag}_logs" / f"{tag}_{source}_to_{TARGET}.pretrain.log"
    text = path.read_text(errors="replace")
    train = [(float(e), float(l)) for l, e in LOSS_RE.findall(text)]
    dev = [(float(e), float(l)) for l, e in EVAL_RE.findall(text)]
    assert train and dev, f"no loss lines in {path}"
    return train, dev


def epoch_means(train):
    """Mean step loss per epoch; a step logged at epoch e (e.g. 3.14) belongs to epoch ceil(e)."""
    buckets = defaultdict(list)
    for e, l in train:
        buckets[-int(-e // 1)].append(l)
    xs = sorted(buckets)
    return xs, [sum(buckets[x]) / len(buckets[x]) for x in xs]


def test_mrr(results_dir, tag, source):
    return json.loads((results_dir / tag / f"{source}_to_{TARGET}" / "eval_test.json").read_text())["eval_mrr"]


def mean_train_loss(tag, source):
    """The Trainer's closing 'train_loss': mean loss over all pretraining steps."""
    path = RUNS / f"{tag}_logs" / f"{tag}_{source}_to_{TARGET}.pretrain.log"
    return float(re.findall(r"'train_loss': ([0-9.]+)", path.read_text(errors="replace"))[-1])


def fresh_sheet(wb, name, index):
    if name in wb.sheetnames:
        wb.remove(wb[name])
    return wb.create_sheet(name, index)


def write_row(ws, row, values, styles):
    for col, (v, st) in enumerate(zip(values, styles), start=1):
        ws.cell(row, col, v)._style = st


def add_loss_sheets(wb):
    main = wb["Original queries"]
    title, head, text, num = (main["A1"]._style, main["A4"]._style, main["B5"]._style, main["L5"]._style)
    formula = main["G5"]._style  # black, 4 decimals

    # One row per run and source.
    ws = fresh_sheet(wb, "Pretrain losses", 2)
    ws.cell(1, 1, "Pretrain losses per run and source (seed 42, target S1000D_spec_sample)")._style = title
    ws.cell(2, 1, "Blue = values read from runs/xml/<tag>_logs/*.pretrain.log and eval_test.json; "
                  "black = formulas. Dev = the source's pretrain dev split.")
    cols = ["Queries", "Pretrain epochs", "Setup", "Run tag", "Source", "Final step loss",
            "Mean train loss (whole run)", "Final dev loss", "Min dev loss", "Epoch of min dev loss",
            "Final − min dev loss", "Test MRR"]
    write_row(ws, 4, cols, [head] * len(cols))
    r = 5
    for name, cfg in SHEETS.items():
        sheet = wb[name]
        for row, tag in cfg["rows"].items():
            if tag is None:
                continue
            for src in SOURCES:
                train, dev = parse_log(tag, src)
                e_min, l_min = min(dev, key=lambda p: p[1])
                write_row(ws, r, [name.split()[0].lower(), sheet.cell(row, 1).value, sheet.cell(row, 2).value, tag,
                                  src, train[-1][1], mean_train_loss(tag, src), dev[-1][1], l_min, e_min,
                                  f"=H{r}-I{r}", test_mrr(cfg["results"], tag, src)],
                          [text] * 5 + [num] * 5 + [formula, num])
                ws.cell(r, 10).number_format = "0.00"
                r += 1
    for col, w in zip("ABCDEFGHIJKL", [11, 10, 30, 30, 15, 11, 13, 11, 11, 11, 11, 10]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[4].height = 42
    ws.freeze_panes = "A5"

    # The 100-epoch curves shown in the figure's top row.
    ws = fresh_sheet(wb, "Dev loss curves", 3)
    ws.cell(1, 1, "Pretrain loss curves, 100-epoch run (longpretrain100_s42, original queries)")._style = title
    ws.cell(2, 1, "Left: mean step loss per epoch. Right: dev loss at each evaluation (every 5 steps; "
                  "steps per epoch differ by source, so evaluation epochs differ too).")
    write_row(ws, 4, ["Epoch"] + [f"Train loss: {s}" for s in SOURCES], [head] * 4)
    curves = {s: parse_log("longpretrain100_s42", s) for s in SOURCES}
    means = {s: dict(zip(*epoch_means(curves[s][0]))) for s in SOURCES}
    for i, ep in enumerate(range(1, 101)):
        write_row(ws, 5 + i, [ep] + [means[s].get(ep) for s in SOURCES], [text] + [num] * 3)
    for j, src in enumerate(SOURCES):
        c = 6 + 3 * j
        ws.cell(4, c, f"Eval epoch: {src}")._style = head
        ws.cell(4, c + 1, f"Dev loss: {src}")._style = head
        for i, (e, l) in enumerate(curves[src][1]):
            ws.cell(5 + i, c, e)._style = num
            ws.cell(5 + i, c + 1, l)._style = num
            ws.cell(5 + i, c).number_format = "0.00"
    for col in range(1, 14):
        ws.column_dimensions[get_column_letter(col)].width = 4 if col in (5, 8, 11) else 12
    ws.row_dimensions[4].height = 42
    ws.freeze_panes = "A5"

    # The figure, with findings computed from the same logs.
    ws = fresh_sheet(wb, "Figure", 2)
    ws.cell(1, 1, "Pretraining overfits after ~15-30 epochs")._style = title
    mins = ", ".join(f"{s} ep {min(curves[s][1], key=lambda p: p[1])[0]:.0f}" for s in SOURCES)
    lines = [
        "Top: 100-epoch run. Train loss (light) keeps falling to ~0.005, but dev loss reaches its minimum "
        f"({mins}) and then rises and plateaus: the model memorises the pretraining data.",
        "Bottom left: dev loss at the end of pretraining for each epoch setting (noisy: a single final evaluation).",
        "Bottom right: mean test MRR after finetuning. Anonymized queries drop from ~0.87 (5-10 ep) to ~0.78 (≥30 ep); "
        "original queries vary less consistently.",
        "The 100-epoch runs were trained without activation checkpointing (GRAD_CKPT=0), unlike the 5/10/30/50 sweep. "
        "Source data: sheets 'Pretrain losses' and 'Dev loss curves'. Image: results/figures/pretrain_overfitting.png",
    ]
    for i, line in enumerate(lines):
        ws.cell(2 + i, 1, line)
    img = Image(str(OUT / "pretrain_overfitting.png"))
    img.width, img.height = img.width * 0.55, img.height * 0.55
    ws.add_image(img, "A7")


def update_xlsx():
    wb = openpyxl.load_workbook(XLSX)
    for name, cfg in SHEETS.items():
        ws = wb[name]
        # Columns L-N hold the final step loss; dev loss goes right after them (O-Q).
        first = 15
        if ws.cell(4, first).value != f"Final pretrain dev loss: {SOURCES[0]}":
            widths = {k: v.width for k, v in ws.column_dimensions.items()}
            ws.insert_cols(first, len(SOURCES))
            for col in range(first, ws.max_column + 1):  # insert_cols does not move column widths
                old = get_column_letter(col - len(SOURCES))
                ws.column_dimensions[get_column_letter(col)].width = widths.get(old, 13) if col >= first + len(SOURCES) else 13
        for i, src in enumerate(SOURCES):
            col = first + i
            head = ws.cell(4, col, f"Final pretrain dev loss: {src}")
            head._style = ws.cell(4, 14)._style
            for row, tag in cfg["rows"].items():
                cell = ws.cell(row, col, parse_log(tag, src)[1][-1][1] if tag else None)
                cell._style = ws.cell(row, 14)._style
    add_loss_sheets(wb)
    notes = wb["Notes"]
    rows = [
        ("Final pretrain dev loss", "last {'eval_loss': ...} line of runs/xml/<tag>_logs/*.pretrain.log: loss on the "
         "source's pretrain dev split at the end of pretraining. It rises after ~15-30 epochs while the step loss "
         "keeps falling (overfitting)"),
        ("Mean train loss", "the Trainer's closing 'train_loss' line: mean step loss over the whole run (sheet "
         "'Pretrain losses'), not the final step"),
        ("Min dev loss", "lowest eval_loss in the run and the (fractional) epoch it occurred at (sheet 'Pretrain losses')"),
        ("Figure", "results/figures/pretrain_overfitting.{png,pdf}, embedded in sheet 'Figure'; curve data in sheet "
         "'Dev loss curves'"),
        ("Rebuild", "losses and the figure: .venv/bin/python viz/plot_pretrain_overfitting.py (then recalc "
         "for cached formula values)"),
    ]
    key_st, val_st = notes.cell(2, 1)._style, notes.cell(2, 2)._style
    keys = [notes.cell(r, 1).value for r in range(1, notes.max_row + 1)]
    for key, detail in rows:
        if key in keys:
            r = keys.index(key) + 1
        else:
            r = keys.index("Rebuild") + 1  # new rows go just above Rebuild
            notes.insert_rows(r)
            keys.insert(r - 1, key)
        notes.cell(r, 1, key)._style = key_st
        notes.cell(r, 2, detail)._style = val_st
    wb.save(XLSX)


def style(ax):
    ax.grid(True, which="major", color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)


def plot():
    fig = plt.figure(figsize=(7.2, 5.2))
    gs = fig.add_gridspec(2, 6, height_ratios=[1, 1], hspace=0.55, wspace=1.4)

    # Top: learning curves of the 100-epoch run (original queries; pretraining never sees the queries).
    top = [fig.add_subplot(gs[0, 2 * i:2 * i + 2]) for i in range(3)]
    for ax, src in zip(top, SOURCES):
        train, dev = parse_log("longpretrain100_s42", src)
        xs, ys = epoch_means(train)
        c = COLORS[src]
        ax.plot(xs, ys, color=c, lw=1.2, alpha=0.45, label="train (epoch mean)")
        ax.plot(*zip(*dev), color=c, lw=2, label="dev")
        e_min, l_min = min(dev, key=lambda p: p[1])
        ax.axvline(e_min, color=MUTED, lw=0.8, ls=(0, (3, 3)))
        ax.plot(e_min, l_min, marker=MARKERS[src], ms=7, color=c, mec="white", mew=1.5, zorder=3)
        ax.text(e_min + 2, 1.4e-3, f"dev min: ep {e_min:.0f}", fontsize=7.5, color=INK, va="bottom")
        ax.set_yscale("log")
        ax.set_ylim(1e-3, 15)
        ax.set_xlim(0, 100)
        ax.set_title(src, fontsize=9, color=INK, loc="left")
        ax.set_xlabel("pretrain epoch")
        style(ax)
    top[0].set_ylabel("pretrain loss (log)")
    proxies = [plt.Line2D([], [], color=MUTED, lw=1.2, alpha=0.45), plt.Line2D([], [], color=MUTED, lw=2)]
    fig.legend(proxies, ["train loss (epoch mean)", "dev loss"], frameon=False, fontsize=7.5, ncol=2,
               loc="lower center", bbox_to_anchor=(0.5, 0.95))

    # Bottom left: final dev loss vs pretrain epochs; bottom right: downstream test MRR.
    ax_l = fig.add_subplot(gs[1, 0:3])
    ax_r = fig.add_subplot(gs[1, 3:6])
    orig, anon = SHEETS["Original queries"], SHEETS["Anonymized queries"]
    mean = {"original": [], "anonymized": []}
    for src in SOURCES:
        dev_final = [parse_log(orig["tag"](ep), src)[1][-1][1] for ep in EPOCHS]
        ax_l.plot(EPOCHS, dev_final, color=COLORS[src], lw=2, marker=MARKERS[src], ms=6, mec="white", mew=1, label=src)
    for ep in EPOCHS:
        for key, cfg in (("original", orig), ("anonymized", anon)):
            mean[key].append(sum(test_mrr(cfg["results"], cfg["tag"](ep), s) for s in SOURCES) / len(SOURCES))
    ax_r.plot(EPOCHS, mean["original"], color=INK, lw=2, marker="o", ms=6, mec="white", mew=1, label="original queries")
    ax_r.plot(EPOCHS, mean["anonymized"], color=MUTED, lw=2, ls=(0, (4, 2)), marker="D", ms=5.5, mec="white", mew=1,
              label="anonymized queries")
    ax_l.set_ylabel("final pretrain dev loss")
    ax_l.set_title("Dev loss at end of pretraining", fontsize=9, color=INK, loc="left")
    ax_l.legend(frameon=False, fontsize=7.5)
    ax_r.set_ylabel("mean test MRR (3 sources)")
    ax_r.set_title("Downstream test MRR after finetuning", fontsize=9, color=INK, loc="left")
    ax_r.legend(frameon=False, fontsize=7.5, loc="center right")
    for ax in (ax_l, ax_r):
        ax.set_xscale("log")
        ax.xaxis.set_major_locator(FixedLocator(EPOCHS))
        ax.xaxis.set_minor_locator(NullLocator())
        ax.set_xticklabels([str(e) for e in EPOCHS])
        ax.set_xlabel("pretrain epochs")
        style(ax)

    OUT.mkdir(parents=True, exist_ok=True)
    for ext in ("pdf", "png"):
        fig.savefig(OUT / f"pretrain_overfitting.{ext}", dpi=200, bbox_inches="tight")


if __name__ == "__main__":
    plot()
    update_xlsx()
